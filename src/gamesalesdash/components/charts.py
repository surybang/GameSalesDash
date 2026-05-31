"""Création des graphiques : barres, courbes et camemberts."""

from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.worksheet.worksheet import Worksheet

from gamesalesdash.config import PIE_COLORS


def create_bar_chart(
    sheet: Worksheet,
    chart_title: str,
    data_range: dict,
    category_range: dict,
    chart_position: str,
    style: int = 11,
    source_sheet: Worksheet | None = None,
) -> None:
    """Crée et insère un graphique à barres horizontales.

    Args:
        sheet: Feuille où le graphique est inséré.
        chart_title: Titre affiché du graphique.
        data_range: Plage des données, clés min_col, min_row, max_row, max_col.
        category_range: Plage des étiquettes, clés min_col, min_row, max_row.
        chart_position: Cellule d'ancrage du graphique (ex : "A23").
        style: Style numérique Excel du graphique.
        source_sheet: Feuille source des données si différente de sheet.
            Si None, sheet est utilisée comme source.
    """
    src = source_sheet or sheet
    chart = BarChart()
    chart.type = "bar"
    chart.style = style
    chart.title = chart_title
    chart.y_axis.title = "Nb jeux vendus (en millions)"
    _attach_data(chart, src, data_range, category_range)
    sheet.add_chart(chart, chart_position)


def create_line_chart(
    sheet: Worksheet,
    chart_title: str,
    data_range: dict,
    category_range: dict,
    chart_position: str,
    y_axis_title: str = "",
    style: int = 11,
    source_sheet: Worksheet | None = None,
) -> None:
    """Crée et insère un graphique en courbes.

    Args:
        sheet: Feuille où le graphique est inséré.
        chart_title: Titre affiché du graphique.
        data_range: Plage des données, clés min_col, min_row, max_row, max_col.
        category_range: Plage des étiquettes, clés min_col, min_row, max_row.
        chart_position: Cellule d'ancrage du graphique (ex : "A1").
        y_axis_title: Libellé de l'axe des ordonnées.
        style: Style numérique Excel du graphique.
        source_sheet: Feuille source des données si différente de sheet.
            Si None, sheet est utilisée comme source.
    """
    src = source_sheet or sheet
    chart = LineChart()
    chart.style = style
    chart.title = chart_title
    chart.y_axis.title = y_axis_title
    _attach_data(chart, src, data_range, category_range)
    sheet.add_chart(chart, chart_position)


def create_pie_chart(
    sheet: Worksheet,
    chart_title: str,
    data_range: dict,
    category_range: dict,
    chart_position: str,
    source_sheet: Worksheet | None = None,
) -> None:
    """Crée et insère un graphique en camembert avec couleurs du dashboard.

    Les tranches utilisent la palette PIE_COLORS définie dans config,
    dérivée des couleurs bleu/vert du dashboard.

    Args:
        sheet: Feuille où le graphique est inséré.
        chart_title: Titre affiché du graphique.
        data_range: Plage des données, clés min_col, min_row, max_row, max_col.
        category_range: Plage des étiquettes, clés min_col, min_row, max_row.
        chart_position: Cellule d'ancrage du graphique (ex : "S1").
        source_sheet: Feuille source des données si différente de sheet.
            Si None, sheet est utilisée comme source.
    """
    src = source_sheet or sheet
    chart = PieChart()
    chart.title = chart_title
    _attach_data(chart, src, data_range, category_range)

    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    chart.dataLabels.showVal = False

    # Couleurs personnalisées par tranche
    num_slices = data_range["max_row"] - data_range["min_row"]
    for i, color in enumerate(PIE_COLORS[:num_slices]):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = color
        chart.series[0].dPt.append(pt)

    sheet.add_chart(chart, chart_position)


def _attach_data(
    chart,
    sheet: Worksheet,
    data_range: dict,
    category_range: dict,
) -> None:
    """Attache les données et les catégories à un graphique.

    Args:
        chart: Graphique openpyxl en cours de construction.
        sheet: Feuille source des données.
        data_range: Plage des données, clés min_col, min_row, max_row, max_col.
        category_range: Plage des étiquettes, clés min_col, min_row, max_row.
    """
    data = Reference(
        sheet,
        min_col=data_range["min_col"],
        min_row=data_range["min_row"],
        max_row=data_range["max_row"],
        max_col=data_range["max_col"],
    )
    cats = Reference(
        sheet,
        min_col=category_range["min_col"],
        min_row=category_range["min_row"],
        max_row=category_range["max_row"],
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
